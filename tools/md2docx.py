#!/usr/bin/env python3
"""Markdown -> DOCX (A4, поля 2 см, Times New Roman, границы таблиц,
разрыв страницы перед каждым заголовком H1 кроме первого).

Использование:  python3 tools/md2docx.py <src.md> <dst.docx> [land] [size]
  land  — ландшафтная ориентация (необязательно)
  size  — размер шрифта основного текста в pt (по умолчанию 12)

Зависимости: pip install pypandoc-binary python-docx
"""
import os
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def ref_doc(path, landscape=False, size=12):
    d = Document()
    s = d.sections[0]
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    else:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2)
    st = d.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    d.save(path)


def main():
    src, dst = sys.argv[1], sys.argv[2]
    landscape = len(sys.argv) > 3 and sys.argv[3] == 'land'
    size = int(sys.argv[4]) if len(sys.argv) > 4 else 12
    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)

    import pypandoc
    with tempfile.TemporaryDirectory() as td:
        ref = os.path.join(td, 'ref.docx')
        ref_doc(ref, landscape, size)
        pypandoc.convert_file(
            src, 'docx', format='markdown', outputfile=dst,
            extra_args=['--reference-doc=' + ref, '--wrap=none'],
        )

    d = Document(dst)
    for t in d.tables:
        t.style = 'Table Grid'
    first = True
    for p in d.paragraphs:
        if p.style.name == 'Heading 1':
            if first:
                first = False
            else:
                p.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)
    d.save(dst)


if __name__ == '__main__':
    main()
